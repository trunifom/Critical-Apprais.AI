"""Table reader for CSV, TSV and XLSX exports (plan chapter 25.5).

Column names differ between tools (Rayyan, Covidence, Scopus, own tables), so headers are matched
to internal columns through an alias table (case, spaces, underscores and hyphens are ignored).
The mapping that was used is returned in ``ReadResult.column_map`` so it can be stored in
``project.yaml`` and the import repeated; ``mapping`` overrides it (``--map abstract=<column>``).
Every source column that has no internal column is kept in ``RawRecord.extra`` as ``col_<header>``.

The alias table is a starting point: real Rayyan/Covidence exports must be checked before such
imports are promised to users (plan chapter 25.5).
"""

from __future__ import annotations

import csv
import io
import logging
import re
from collections import Counter
from collections.abc import Iterable
from datetime import date, datetime
from pathlib import Path
from typing import Any

from crapai.errors import ImportFailed
from crapai.io.readers.base import RawRecord, ReadResult, read_text_file
from crapai.io.readers.detect import SourceFormat, detect_format

logger = logging.getLogger(__name__)

# Internal column -> accepted header names in priority order (normalised, see _normalise).
ALIASES: dict[str, tuple[str, ...]] = {
    "title": ("title", "titel", "ti", "t1", "articletitle", "documenttitle", "primarytitle"),
    "abstract": ("abstract", "zusammenfassung", "ab", "n2", "abstractnote", "summary"),
    "authors": ("authors", "autoren", "author", "au", "a1", "authorfullnames"),
    "year": ("year", "jahr", "py", "publicationyear", "publishedyear", "date"),
    "journal": (
        "journal",
        "zeitschrift",
        "sourcetitle",
        "publicationtitle",
        "jo",
        "jf",
        "t2",
        "journalname",
    ),
    "doi": ("doi", "do", "digitalobjectidentifier"),
    "pmid": ("pmid", "pubmedid"),
    "keywords": ("keywords", "schlagwörter", "schlagworter", "kw", "authorkeywords"),
    "url": ("url", "link", "ur"),
}
# Columns that must stay text: Excel and CSV round trips turn them into floats or lose zeros.
TEXT_COLUMNS = frozenset({"pmid", "doi"})

YEAR = re.compile(r"\b(1[4-9]\d{2}|20\d{2}|2100)\b")
FLOAT_INTEGER = re.compile(r"^(\d+)\.0+$")
LARGE_SHEET_ROWS = 200_000


def _normalise(header: str) -> str:
    """Normalise a header for alias matching: lower case, no spaces, underscores or hyphens."""
    return re.sub(r"[\s_\-]+", "", header.strip().lower())


def read_table(
    path: Path,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | None = None,
    mapping: dict[str, str] | None = None,
) -> ReadResult:
    """Read a CSV, TSV or XLSX file into records.

    Args:
        path: The table file.
        encoding: Force a text encoding (CSV/TSV); default is the chain of plan chapter 25.9.
        delimiter: Force the field separator; default is detected.
        sheet: Worksheet name for XLSX; default is the first visible sheet.
        mapping: Explicit ``{internal column: source header}`` pairs (``--map``); they win over
            the alias table.

    Raises:
        ImportFailed: E101 (not a table), E102 (no data rows), E103 (undecodable), E104 (a
            requested column or a title/abstract column cannot be found).
    """
    detected = detect_format(path)
    if detected.format is SourceFormat.XLSX:
        rows, options, notes = _read_xlsx_rows(path, sheet)
        used_encoding = None
    elif detected.format in (SourceFormat.CSV, SourceFormat.TSV):
        chosen = _normalise_delimiter(delimiter) if delimiter else detected.delimiter or ","
        text, used_encoding = read_text_file(path, encoding)
        rows = _read_csv_rows(text, chosen)
        options = {"delimiter": chosen}
        notes = []
        replaced = text.count("�")
        if replaced:
            notes.append(f"{replaced} replacement character(s) in the text (damaged in the source)")
    else:
        raise ImportFailed(
            f"{path.name} is not a table (detected: {detected.format.value})",
            code="E101",
            hint="Use the reader for that format, or export the data as CSV or XLSX.",
            details={"path": str(path)},
        )
    result = _rows_to_records(rows, mapping or {}, detected.format)
    result.encoding = used_encoding
    result.options.update(options)
    result.notes[:0] = notes
    logger.info("Read %d table row(s) from %s", len(result.records), path.name)
    return result


# --- raw rows -----------------------------------------------------------------------------


DELIMITER_NAMES = {"\\t": "\t", "tab": "\t", "comma": ",", "semicolon": ";", "pipe": "|"}


def _normalise_delimiter(value: str) -> str:
    """One separator character from ``--delimiter``; accepts ``\\t``, ``tab``, ``semicolon`` ...

    Raises:
        ImportFailed: (E104) if the value is not exactly one character after the names are
            resolved; ``csv`` would fail with a raw ``TypeError`` otherwise.
    """
    resolved = DELIMITER_NAMES.get(value.strip().lower(), value)
    if len(resolved) != 1:
        raise ImportFailed(
            f"The field separator must be one character, got '{value}'",
            code="E104",
            hint="Use , or ; or | or a tab (write: tab), for example --delimiter semicolon.",
            details={"delimiter": value},
        )
    return resolved


def _read_csv_rows(text: str, delimiter: str) -> list[list[str]]:
    csv.field_size_limit(max(csv.field_size_limit(), 10_000_000))
    return [
        row
        for row in csv.reader(io.StringIO(text), delimiter=delimiter)
        if any(cell.strip() for cell in row)
    ]


def _cell_text(value: Any) -> str:
    """Text of one Excel cell: no exponent notation for integers, ISO dates, no ``None``."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        midnight = value.time() == datetime.min.time()
        return value.date().isoformat() if midnight else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _read_xlsx_rows(
    path: Path, sheet: str | None
) -> tuple[list[list[str]], dict[str, Any], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - the extra is normally installed
        raise ImportFailed(
            "Reading Excel files needs the 'import' extra (openpyxl)",
            code="E101",
            hint='Install it with: pip install -e ".[import]"',
        ) from exc
    notes: list[str] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for damaged files
        raise ImportFailed(
            f"{path.name} cannot be opened as an Excel workbook ({type(exc).__name__})",
            code="E101",
            hint="Open and save the file again in Excel, or export it as CSV.",
            details={"path": str(path)},
        ) from exc
    try:
        visible = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
        if sheet is not None:
            chosen = next((ws for ws in visible if ws.title == sheet), None)
            if chosen is None:
                raise ImportFailed(
                    f"Worksheet '{sheet}' not found in {path.name} "
                    f"(sheets: {', '.join(ws.title for ws in visible)})",
                    code="E101",
                    hint="Available sheets: " + ", ".join(ws.title for ws in visible),
                    details={"path": str(path)},
                )
        elif visible:
            chosen = visible[0]
            if len(visible) > 1:
                notes.append(f"workbook has {len(visible)} sheets; read '{chosen.title}'")
        else:
            raise ImportFailed(f"{path.name} has no visible worksheet", code="E102")
        try:
            rows = [
                [_cell_text(cell) for cell in row]
                for row in chosen.iter_rows(values_only=True)
                if any(cell is not None and str(cell).strip() for cell in row)
            ]
        except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for damaged sheets
            logger.error("Cannot read the rows of %s (%s)", path.name, type(exc).__name__)
            raise ImportFailed(
                f"{path.name}: the worksheet is damaged ({type(exc).__name__})",
                code="E101",
                hint="Open and save the file again in Excel, or export it as CSV.",
                details={"path": str(path)},
            ) from exc
        if len(rows) > LARGE_SHEET_ROWS:
            notes.append(f"large sheet ({len(rows)} rows)")
        return rows, {"sheet": chosen.title}, notes
    finally:
        workbook.close()


# --- header and mapping -------------------------------------------------------------------


def unique_headers(raw: Iterable[str]) -> list[str]:
    """Trim headers, name empty ones ``col_N`` and number duplicates (``title``, ``title_2``)."""
    headers: list[str] = []
    seen: Counter[str] = Counter()
    for index, name in enumerate(raw, start=1):
        clean = name.strip() or f"col_{index}"
        seen[clean] += 1
        headers.append(clean if seen[clean] == 1 else f"{clean}_{seen[clean]}")
    return headers


def resolve_columns(
    headers: list[str], mapping: dict[str, str]
) -> tuple[dict[str, str], list[str]]:
    """Choose the source header for every internal column.

    Returns:
        ``(column_map, notes)`` with ``column_map`` = internal column -> source header.

    Raises:
        ImportFailed: (E104) if ``mapping`` names an unknown internal column or header.
    """
    notes: list[str] = []
    chosen: dict[str, str] = {}
    for internal, header in mapping.items():
        if internal not in ALIASES:
            raise ImportFailed(
                f"Unknown target column '{internal}' in the mapping (valid: {', '.join(ALIASES)})",
                code="E104",
                hint="Valid targets: " + ", ".join(ALIASES),
            )
        if header not in headers:
            raise ImportFailed(
                f"Column '{header}' not found in the file (columns: {', '.join(headers)})",
                code="E104",
                hint="Available columns: " + ", ".join(headers),
                details={"column": header},
            )
        chosen[internal] = header
    normalised = {header: _normalise(header) for header in headers}
    for internal, aliases in ALIASES.items():
        if internal in chosen:
            continue
        matches = [header for alias in aliases for header in headers if normalised[header] == alias]
        if matches:
            chosen[internal] = matches[0]
            others = [h for h in dict.fromkeys(matches) if h != matches[0]]
            if others:
                notes.append(
                    f"column '{matches[0]}' was used for {internal}; also matching: "
                    + ", ".join(others)
                    + f" (use --map {internal}=<column> to change)"
                )
    return chosen, notes


def _rows_to_records(
    rows: list[list[str]], mapping: dict[str, str], source_format: SourceFormat
) -> ReadResult:
    if len(rows) < 2:
        raise ImportFailed(
            "The table has no data rows",
            code="E102",
            hint="Check that the first row holds the column names and data follows.",
        )
    headers = unique_headers(rows[0])
    column_map, notes = resolve_columns(headers, mapping)
    if "title" not in column_map and "abstract" not in column_map:
        raise ImportFailed(
            "Neither a title nor an abstract column could be identified",
            code="E104",
            hint="Use --map title=<column> and --map abstract=<column>. Columns: "
            + ", ".join(headers),
            details={"columns": headers},
        )
    if "abstract" not in column_map:
        notes.append("no abstract column found; use --map abstract=<column>")
    mapped_headers = set(column_map.values())
    position = {header: index for index, header in enumerate(headers)}
    records: list[RawRecord] = []
    for number, row in enumerate(rows[1:], start=1):
        cells = row + [""] * (len(headers) - len(row))
        record = RawRecord(number, {"record_type": "other"})
        for internal, header in column_map.items():
            value = _clean_cell(internal, cells[position[header]])
            if internal == "year":
                year = YEAR.search(value)
                if year:
                    record.fields["year"] = int(year.group(1))
                elif value:
                    # Not a 4-digit year ("n.d.", "in press", "forthcoming", ...): nothing is
                    # lost, the raw cell survives under the same key an unmapped column would
                    # use (it is otherwise unused here, since "year"'s header is itself mapped).
                    record.extra[f"col_{header}"] = value
            elif value:
                record.fields[internal] = value
        for header in headers:
            if header not in mapped_headers and cells[position[header]].strip():
                record.extra[f"col_{header}"] = cells[position[header]].strip()
        if len(row) > len(headers):
            record.notes.append(f"{len(row) - len(headers)} cell(s) beyond the header were ignored")
            record.extra["col_overflow"] = [c for c in row[len(headers) :] if c.strip()]
        records.append(record)
    return ReadResult(source_format, records, notes=notes, column_map=column_map)


def _clean_cell(internal: str, value: str) -> str:
    text = value.strip()
    if internal in TEXT_COLUMNS:
        match = FLOAT_INTEGER.match(text)
        if match:  # "32559806.0" from a spreadsheet round trip
            return match.group(1)
    return text
