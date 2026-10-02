"""Export of records (CSV, XLSX, RIS) and of the PRISMA flow, and `crapai export` (review task)."""

from __future__ import annotations

import csv
import datetime as dt
import json
from pathlib import Path
from typing import Any

import pytest
from results_helpers import project_with_run
from typer.testing import CliRunner

from crapai.cli import app
from crapai.errors import ConfigError, ImportFailed, StorageError
from crapai.io.readers.bibtex import parse_bibtex_text
from crapai.io.readers.nbib import parse_nbib_text
from crapai.io.readers.ris import parse_ris_text
from crapai.io.records_store import RECORD_COLUMNS, Record
from crapai.io.writers.bibtex import record_to_bibtex, write_bibtex
from crapai.io.writers.nbib import record_to_nbib, write_nbib
from crapai.io.writers.ris import record_to_ris, write_ris
from crapai.io.writers.tables import neutralise_formula, write_csv, write_xlsx
from crapai.project import atomic
from crapai.project.workspace import Workspace
from crapai.services.dedup import dedup_project
from crapai.services.export import (
    export_flow,
    export_prospero,
    export_records,
    export_report,
    select_records,
)
from crapai.services.importing import ImportRequest, import_source
from crapai.services.preflight import check_project
from crapai.services.project import create_project

runner = CliRunner()
NOW = dt.datetime(2026, 10, 1, 9, 30, 15)
ABSTRACT = "The trial enrolled adults and measured outcomes over twelve months. " * 5


def record(**values: Any) -> Record:
    base: dict[str, Any] = {
        "study_uid": "uid-1",
        "source_label": "S",
        "source_file": "f.ris",
        "source_row": 1,
        "source_format": "ris",
        "title": "A study of exercise",
    }
    base.update(values)
    return Record.model_validate(base)


@pytest.fixture
def project(tmp_path: Path) -> Workspace:
    workspace = create_project(tmp_path / "p", template="demo")
    text = (
        "TY  - JOUR\nTI  - Effects of exercise on mood\nAU  - Doe, Jane\nAU  - Roe, Richard\n"
        "PY  - 2020\nJO  - Journal of Tests\nVL  - 7\nIS  - 2\n"
        "SP  - 10\nEP  - 20\nDO  - 10.1000/a\n"
        f"KW  - sport\nKW  - mood\nAB  - {ABSTRACT}\nER  - \n"
        "TY  - JOUR\nTI  - Effects of exercise on mood\nDO  - 10.1000/a\nER  - \n"
        'TY  - JOUR\nTI  - =HYPERLINK("http://evil","x")\nAB  - text with umlaut äöü\nER  - \n'
    )
    source = tmp_path / "s.ris"
    source.write_text(text, encoding="utf-8")
    import_source(workspace, ImportRequest(source, label="S"))
    check_project(workspace)
    return workspace


# --- the writers -------------------------------------------------------------------------------


@pytest.mark.parametrize("value", ["=1+1", "+1", "-1", "@cmd", "\tx", "\rx"])
def test_formula_starts_are_neutralised(value: str) -> None:
    assert neutralise_formula(value) == "'" + value


@pytest.mark.parametrize("value", ["", "plain", "a=b", "10.1000/x", "Über"])
def test_ordinary_text_is_untouched(value: str) -> None:
    assert neutralise_formula(value) == value


def test_csv_has_bom_crlf_and_quotes_where_needed(tmp_path: Path) -> None:
    path = tmp_path / "t.csv"
    write_csv(path, ["a", "b"], [["x,y", 'he said "hi"'], ["Größe", "=2+2"]])
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and b"\r\n" in raw
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines()))
    assert rows == [["a", "b"], ["x,y", 'he said "hi"'], ["Größe", "'=2+2"]]


def test_csv_separator_and_raw_mode(tmp_path: Path) -> None:
    path = tmp_path / "t.csv"
    write_csv(path, ["a", "b"], [["1", "=2"]], delimiter=";", guard_formulas=False)
    assert path.read_bytes().decode("utf-8-sig") == "a;b\r\n1;=2\r\n"


@pytest.mark.parametrize("delimiter", ["", "::"])
def test_csv_separator_must_be_one_character(tmp_path: Path, delimiter: str) -> None:
    with pytest.raises(ConfigError) as info:
        write_csv(tmp_path / "t.csv", ["a"], [], delimiter=delimiter)
    assert info.value.code == "E203"


def test_xlsx_is_readable_frozen_and_safe(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "t.xlsx"
    write_xlsx(path, ["title", "note"], [["=SUM(A1)", "ok\x07bell"], ["Größe", "x" * 40_000]])
    sheet = openpyxl.load_workbook(path).active
    assert sheet.freeze_panes == "A2" and sheet["A1"].font.bold
    assert sheet["A2"].value == "'=SUM(A1)" and sheet["B2"].value == "okbell"
    assert sheet["A3"].value == "Größe" and len(sheet["B3"].value) == 32_767


def test_xlsx_without_openpyxl_is_e203(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import sys

    monkeypatch.setitem(sys.modules, "openpyxl", None)
    with pytest.raises(ConfigError) as info:
        write_xlsx(tmp_path / "t.xlsx", ["a"], [])
    assert info.value.code == "E203" and "openpyxl" in info.value.user_message


def test_ris_entry_carries_the_bibliographic_fields() -> None:
    entry = record_to_ris(
        record(
            authors="Doe, Jane; Roe, Richard", year=2020, journal="J Tests", volume="7",
            issue="2", pages="10-20", doi="10.1000/a", keywords="a; b", keywords_mesh="m",
            abstract="Line one\nline two", exclusion_reason="DUPLICATE", is_duplicate=True,
            duplicate_of="uid-0", is_retracted=True, record_type="journal_article", pmid="99",
            language="eng",
        )
    )  # fmt: skip
    lines = entry.split("\n")
    assert lines[0] == "TY  - JOUR" and lines[-1] == "ER  - "
    for expected in (
        "ID  - uid-1", "TI  - A study of exercise", "AU  - Doe, Jane", "AU  - Roe, Richard",
        "PY  - 2020", "JO  - J Tests", "VL  - 7", "IS  - 2", "SP  - 10", "EP  - 20",
        "DO  - 10.1000/a", "AN  - 99", "LA  - eng", "KW  - a", "KW  - b", "KW  - m",
        "AB  - Line one line two",
        "N1  - Excluded before screening: DUPLICATE", "N1  - Duplicate of: uid-0",
        "N1  - Retracted publication",
    ):  # fmt: skip
        assert expected in lines


def test_ris_types_round_trip_and_unknown_types_become_generic() -> None:
    assert record_to_ris(record(record_type="book")).startswith("TY  - BOOK")
    assert record_to_ris(record(record_type="conference_paper")).startswith("TY  - CONF")
    assert record_to_ris(record(record_type="other")).startswith("TY  - GEN")


def test_ris_written_by_us_is_read_back_by_our_reader(tmp_path: Path) -> None:
    original = [
        record(study_uid="u1", authors="Doe, Jane; Roe, Richard", year=2020, doi="10.1000/a",
               abstract="An abstract", journal="J", pages="1-5"),
        record(study_uid="u2", title="Second study", record_type="book"),
    ]  # fmt: skip
    path = tmp_path / "out.ris"
    write_ris(path, original)
    back = parse_ris_text(path.read_text(encoding="utf-8")).records
    assert [r.fields["title"] for r in back] == ["A study of exercise", "Second study"]
    assert back[0].fields["authors"] == "Doe, Jane; Roe, Richard"
    assert back[0].fields["year"] == 2020 and back[0].fields["pages"] == "1-5"
    assert back[1].fields["record_type"] == "book"


def test_an_empty_ris_export_is_an_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "e.ris"
    write_ris(path, [])
    assert path.read_text(encoding="utf-8") == ""


def test_bibtex_entry_carries_the_bibliographic_fields() -> None:
    entry = record_to_bibtex(
        record(
            authors="Doe, Jane; Roe, Richard", year=2020, journal="J Tests", volume="7",
            issue="2", pages="10-20", doi="10.1000/a", keywords="a; b", keywords_mesh="m",
            abstract="Line one\nline two", exclusion_reason="DUPLICATE", is_duplicate=True,
            duplicate_of="uid-0", is_retracted=True, record_type="journal_article",
        )
    )  # fmt: skip
    assert entry.startswith("@article{uid-1,")
    for expected in (
        "title = {A study of exercise}", "author = {Doe, Jane and Roe, Richard}",
        "year = {2020}", "journal = {J Tests}", "volume = {7}", "number = {2}",
        "pages = {10-20}", "doi = {10.1000/a}", "keywords = {a; b; m}",
        "abstract = {Line one line two}",
        "note = {Excluded before screening: DUPLICATE; Duplicate of: uid-0; Retracted publication}",
    ):  # fmt: skip
        assert expected in entry
    assert entry.endswith("\n}")


def test_bibtex_types_round_trip_and_unknown_types_become_misc() -> None:
    assert record_to_bibtex(record(record_type="book")).startswith("@book{")
    assert record_to_bibtex(record(record_type="thesis")).startswith("@phdthesis{")
    assert record_to_bibtex(record(record_type="other")).startswith("@misc{")


def test_bibtex_braces_and_backslashes_in_a_value_are_escaped() -> None:
    entry = record_to_bibtex(record(title=r"A {nested} \value"))
    assert r"A \{nested\} \\value" in entry
    # the escaping must not leave the braces unbalanced
    assert entry.count("{") == entry.count("}")


def test_bibtex_written_by_us_is_read_back_by_our_reader(tmp_path: Path) -> None:
    original = [
        record(study_uid="u1", authors="Doe, Jane; Roe, Richard", year=2020, doi="10.1000/a",
               abstract="An abstract", journal="J", pages="1-5"),
        record(study_uid="u2", title="Second study", record_type="book"),
    ]  # fmt: skip
    path = tmp_path / "out.bib"
    write_bibtex(path, original)
    back = parse_bibtex_text(path.read_text(encoding="utf-8")).records
    assert [r.fields["title"] for r in back] == ["A study of exercise", "Second study"]
    assert back[0].fields["authors"] == "Doe, Jane; Roe, Richard"
    assert back[0].fields["year"] == 2020 and back[0].fields["pages"] == "1-5"
    assert back[1].fields["record_type"] == "book"


def test_an_empty_bibtex_export_is_an_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "e.bib"
    write_bibtex(path, [])
    assert path.read_text(encoding="utf-8") == ""


def test_nbib_entry_carries_the_bibliographic_fields() -> None:
    entry = record_to_nbib(
        record(
            pmid="12345", authors="Doe, Jane; Roe, Richard", year=2020, journal="J Tests",
            volume="7", issue="2", pages="10-20", doi="10.1000/a", keywords="a",
            keywords_mesh="m", abstract="Line one\nline two", pmcid="PMC1",
        )
    )  # fmt: skip
    lines = entry.split("\n")
    for expected in (
        "PMID- 12345", "UID - uid-1", "TI  - A study of exercise", "FAU - Doe, Jane",
        "FAU - Roe, Richard", "JT  - J Tests", "TA  - J Tests", "DP  - 2020", "VI  - 7",
        "IP  - 2", "PG  - 10-20", "AID - 10.1000/a [doi]", "PMC - PMC1", "MH  - m",
        "OT  - a", "AB  - Line one line two",
    ):  # fmt: skip
        assert expected in lines


def test_nbib_written_by_us_is_read_back_by_our_reader(tmp_path: Path) -> None:
    original = [
        record(study_uid="u1", pmid="111", authors="Doe, Jane", year=2020, doi="10.1000/a",
               abstract="An abstract", journal="J", pages="1-5"),
        record(study_uid="u2", pmid="222", title="Second study"),
    ]  # fmt: skip
    path = tmp_path / "out.nbib"
    write_nbib(path, original)
    back = parse_nbib_text(path.read_text(encoding="utf-8")).records
    assert [r.fields["title"] for r in back] == ["A study of exercise", "Second study"]
    assert back[0].fields["authors"] == "Doe, Jane"
    assert back[0].fields["doi"] == "10.1000/a" and back[0].fields["pages"] == "1-5"


def test_nbib_export_without_any_pmid_cannot_be_read_back(tmp_path: Path) -> None:
    """A known limitation (documented in the writer): NBIB's own reader uses PMID as its
    defining signal for "this file is MEDLINE", so a file with zero PMIDs fails to be recognised
    at all, even though every record still has its own full, correctly tagged entry."""
    path = tmp_path / "no_pmid.nbib"
    write_nbib(path, [record(title="No PMID here")])
    with pytest.raises(ImportFailed) as info:
        parse_nbib_text(path.read_text(encoding="utf-8"))
    assert info.value.code == "E102"


def test_an_empty_nbib_export_is_an_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "e.nbib"
    write_nbib(path, [])
    assert path.read_text(encoding="utf-8") == ""


# --- the service -------------------------------------------------------------------------------


def test_scopes_select_the_right_records(project: Workspace) -> None:
    from crapai.io.records_store import read_records

    records = read_records(project.records_csv)
    assert len(select_records(records, "all")) == 3
    screenable = select_records(records, "screenable")
    excluded = select_records(records, "excluded")
    assert len(screenable) + len(excluded) == 3 and all(r.exclusion_reason for r in excluded)
    with pytest.raises(ConfigError) as info:
        select_records(records, "some")
    assert info.value.code == "E203"


def test_csv_export_goes_to_the_exports_folder_with_all_columns(project: Workspace) -> None:
    summary = export_records(project, "csv", now=NOW)
    assert summary.path == project.exports_dir / "records-all-20261001-093015.csv"
    assert summary.records == 3 and not summary.used_alternative
    rows = list(csv.reader(summary.path.read_text(encoding="utf-8-sig").splitlines()))
    assert tuple(rows[0]) == RECORD_COLUMNS and len(rows) == 4


def test_the_csv_export_neutralises_a_hostile_title(project: Workspace) -> None:
    summary = export_records(project, "csv", now=NOW)
    text = summary.path.read_text(encoding="utf-8-sig")
    assert "'=HYPERLINK" in text
    raw = export_records(project, "csv", guard_formulas=False, output=project.root / "raw.csv")
    assert ",=HYPERLINK" in raw.path.read_text(
        encoding="utf-8-sig"
    ) or '"=HYPERLINK' in raw.path.read_text(encoding="utf-8-sig")


def test_the_export_does_not_change_the_project(project: Workspace) -> None:
    before = {p: p.read_bytes() for p in project.root.rglob("*") if p.is_file()}
    export_records(project, "csv", now=NOW)
    export_records(project, "ris", now=NOW)
    changed = {p for p, data in before.items() if p.read_bytes() != data}
    assert changed == set()  # only new files appear (under exports/)


def test_scope_screenable_and_excluded(project: Workspace) -> None:
    dedup_project(project)
    check_project(project)
    both = [
        export_records(project, "ris", scope=s, now=NOW).records for s in ("screenable", "excluded")
    ]
    assert sum(both) == 3


def test_xlsx_and_ris_exports(project: Workspace) -> None:
    pytest.importorskip("openpyxl")
    assert export_records(project, "xlsx", now=NOW).path.suffix == ".xlsx"
    ris = export_records(project, "ris", now=NOW)
    assert ris.path.read_text(encoding="utf-8").count("ER  - ") == 3


def test_bibtex_and_nbib_exports(project: Workspace) -> None:
    bibtex = export_records(project, "bibtex", now=NOW)
    assert bibtex.path.suffix == ".bib" and bibtex.path.read_text(encoding="utf-8").count("@") >= 3
    nbib = export_records(project, "nbib", now=NOW)
    assert nbib.path.suffix == ".nbib"


def test_a_fulltext_attachment_is_not_exported_as_its_own_bibliographic_entry(
    project: Workspace,
) -> None:
    """Regression test (found live, end-to-end): a full-text attachment (ADR 0026) legitimately
    shares its DOI/title with the record it belongs to. CSV keeps it (the full data dump), but a
    reference-manager format must not gain a spurious near-duplicate entry for every PDF import."""
    from crapai.io.records_store import read_records, write_records

    anchor_uid = "uid-1"  # the "Effects of exercise on mood" record already in `project`
    attachment = record(
        study_uid="uid-attachment",
        title="DOI: 10.1000/a",
        doi="10.1000/a",
        fulltext_of=anchor_uid,
        has_fulltext=True,
    )
    write_records(project.records_csv, [*read_records(project.records_csv), attachment])

    bibtex = export_records(project, "bibtex", now=NOW)
    assert "DOI: 10.1000/a" not in bibtex.path.read_text(encoding="utf-8")
    nbib = export_records(project, "nbib", now=NOW)
    assert "DOI: 10.1000/a" not in nbib.path.read_text(encoding="utf-8")
    ris = export_records(project, "ris", now=NOW)
    assert "DOI: 10.1000/a" not in ris.path.read_text(encoding="utf-8")

    csv_export = export_records(project, "csv", now=NOW)
    assert "DOI: 10.1000/a" in csv_export.path.read_text(encoding="utf-8-sig")


def test_an_explicit_output_path_is_used(project: Workspace, tmp_path: Path) -> None:
    target = tmp_path / "mine.csv"
    assert export_records(project, "csv", output=target).path == target and target.exists()


@pytest.mark.parametrize("kind", ["folder", "missing_parent"])
def test_bad_targets_are_e203(project: Workspace, tmp_path: Path, kind: str) -> None:
    target = tmp_path if kind == "folder" else tmp_path / "nope" / "x.csv"
    with pytest.raises(ConfigError) as info:
        export_records(project, "csv", output=target)
    assert info.value.code == "E203"


def test_unknown_format_is_e203(project: Workspace) -> None:
    with pytest.raises(ConfigError) as info:
        export_records(project, "pdf")
    assert info.value.code == "E203" and "csv, xlsx, ris" in info.value.user_message


def test_a_locked_target_gives_an_alternative_file(
    project: Workspace, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "results.csv"
    target.write_text("old", encoding="utf-8")
    real = atomic.os.replace

    def locked(src: object, dst: object) -> None:
        if Path(str(dst)) == target:
            raise PermissionError("open in Excel")
        real(src, dst)  # type: ignore[arg-type]

    monkeypatch.setattr(atomic.os, "replace", locked)
    monkeypatch.setattr(atomic, "DEFAULT_DELAY_S", 0.0)
    summary = export_records(project, "csv", output=target)
    assert summary.used_alternative and summary.path != target and summary.requested_path == target
    assert target.read_text(encoding="utf-8") == "old"


def test_the_export_of_a_folder_without_project_is_e404(tmp_path: Path) -> None:
    with pytest.raises(StorageError) as info:
        export_records(Workspace(tmp_path / "nothing"), "csv")
    assert info.value.code == "E404"


def test_flow_export_has_numbers_warnings_and_mode(project: Workspace) -> None:
    summary = export_flow(project, now=NOW)
    assert summary.path == project.exports_dir / "prisma_flow.json" and summary.format == "json"
    document = json.loads(summary.path.read_text(encoding="utf-8"))
    assert document["schema"] == 1 and document["reporting_mode"] == "all_before_screening"
    assert document["flow"]["records_identified_total"] == 3
    assert document["flow"]["duplicates_removed"] == 1 and document["warnings"] == []
    assert document["events"] >= 4 and document["generated_at"].startswith("2026-10-01T")
    assert "ta_included" not in document["flow"]


def test_flow_export_in_the_other_mode_and_with_a_bad_mode(project: Workspace) -> None:
    summary = export_flow(project, mode="between_databases_only", now=NOW)
    assert json.loads(summary.path.read_text(encoding="utf-8"))["reporting_mode"] == (
        "between_databases_only"
    )
    with pytest.raises(ConfigError):
        export_flow(project, mode="sometimes")


def test_flow_export_as_png_or_svg(project: Workspace) -> None:
    pytest.importorskip("matplotlib")
    png = export_flow(project, fmt="png", now=NOW)
    assert png.path.suffix == ".png" and png.path.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
    svg = export_flow(project, fmt="svg", now=NOW)
    assert svg.path.suffix == ".svg" and "Records identified" in svg.path.read_text(
        encoding="utf-8"
    )


def test_flow_export_rejects_an_unknown_format(project: Workspace) -> None:
    with pytest.raises(ConfigError) as info:
        export_flow(project, fmt="gif")
    assert info.value.code == "E203"


def test_docx_report_without_any_run_still_writes_a_report(project: Workspace) -> None:
    pytest.importorskip("docx")
    from docx import Document

    summary = export_report(project, now=NOW)
    assert summary.path.suffix == ".docx" and summary.what == "report"
    text = "\n".join(p.text for p in Document(summary.path).paragraphs)
    assert "No finished screening run yet." in text


def test_docx_report_with_a_finished_run(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    workspace = project_with_run(tmp_path, scenario="S1")
    summary = export_report(workspace, now=NOW)
    text = "\n".join(p.text for p in Document(summary.path).paragraphs)
    assert "No finished screening run yet." not in text
    assert "record(s) screened" in text
    headers = [cell.text for t in Document(summary.path).tables for cell in t.rows[0].cells]
    assert "PRISMA step" in headers and "Decision" in headers


def test_prospero_export_without_metadata_shows_placeholders(project: Workspace) -> None:
    pytest.importorskip("docx")
    from docx import Document

    summary = export_prospero(project, now=NOW)
    assert summary.path.suffix == ".docx" and summary.what == "prospero"
    text = "\n".join(p.text for p in Document(summary.path).paragraphs)
    tables = [cell.text for t in Document(summary.path).tables for row in t.rows for cell in row.cells]
    assert "PROSPERO" in text and "no public submission API" in text
    assert "(please complete -- not tracked by this project)" in tables
    assert "Include Population" in text  # the project fixture's own criteria still show up


def test_prospero_export_fills_in_configured_metadata(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    import yaml
    from docx import Document

    workspace = create_project(tmp_path / "p2", template="demo")
    data = yaml.safe_load(workspace.project_yaml.read_text(encoding="utf-8"))
    data["prospero"] = {
        "anticipated_start_date": "2026-01-01",
        "team_members": ["Jane Doe", "John Roe"],
        "funding": "None",
        "conflicts_of_interest": "None known",
    }
    workspace.project_yaml.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
    summary = export_prospero(workspace, now=NOW)
    tables = [cell.text for t in Document(summary.path).tables for row in t.rows for cell in row.cells]
    assert "2026-01-01" in tables
    assert "Jane Doe; John Roe" in tables
    assert "None" in tables and "None known" in tables


def test_flow_export_carries_stale_warnings(project: Workspace, tmp_path: Path) -> None:
    more = tmp_path / "more.ris"
    more.write_text(
        f"TY  - JOUR\nTI  - Another unrelated study\nAB  - {ABSTRACT}\nER  - \n", encoding="utf-8"
    )
    import_source(project, ImportRequest(more, label="T"))
    document = json.loads(export_flow(project).path.read_text(encoding="utf-8"))
    assert [w["code"] for w in document["warnings"]] == ["STALE_DEDUP"]


# --- the command -------------------------------------------------------------------------------


def test_export_command_csv(project: Workspace) -> None:
    result = runner.invoke(app, ["export", str(project.root), "--lang", "en"])
    assert result.exit_code == 0, result.output
    assert "Exported 3 record(s) (scope: all) as csv" in result.output
    assert len(list(project.exports_dir.glob("records-all-*.csv"))) == 1


def test_export_command_german_and_semicolon(project: Workspace, tmp_path: Path) -> None:
    target = tmp_path / "de.csv"
    result = runner.invoke(
        app,
        [
            "export",
            str(project.root),
            "--delimiter",
            "semicolon",
            "--output",
            str(target),
            "--lang",
            "de",
        ],
    )
    assert result.exit_code == 0 and "exportiert" in result.output
    assert target.read_text(encoding="utf-8-sig").startswith("study_uid;source_label;")


def test_export_command_flow_json(project: Workspace) -> None:
    result = CliRunner().invoke(app, ["export", str(project.root), "--what", "flow", "--json"])
    assert result.exit_code == 0
    data = json.loads(result.stdout)
    assert data["what"] == "flow" and Path(data["path"]).name == "prisma_flow.json"


def test_export_command_flow_as_png(project: Workspace) -> None:
    pytest.importorskip("matplotlib")
    result = CliRunner().invoke(
        app, ["export", str(project.root), "--what", "flow", "--format", "png"]
    )
    assert result.exit_code == 0, result.output
    assert "prisma_flow.png" in result.output


def test_export_command_report(project: Workspace) -> None:
    pytest.importorskip("docx")
    result = CliRunner().invoke(app, ["export", str(project.root), "--what", "report"])
    assert result.exit_code == 0, result.output
    assert ".docx" in result.output


def test_export_command_reports_a_locked_target_with_exit_four(
    project: Workspace, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "busy.csv"
    target.write_text("old", encoding="utf-8")
    real = atomic.os.replace

    def locked(src: object, dst: object) -> None:
        if Path(str(dst)) == target:
            raise PermissionError("open in Excel")
        real(src, dst)  # type: ignore[arg-type]

    monkeypatch.setattr(atomic.os, "replace", locked)
    monkeypatch.setattr(atomic, "DEFAULT_DELAY_S", 0.0)
    result = runner.invoke(
        app, ["export", str(project.root), "--output", str(target), "--lang", "en"]
    )
    assert result.exit_code == 4 and "open in another program" in result.stderr


@pytest.mark.parametrize(
    "args,code",
    [(["--what", "x"], "E203"), (["--format", "pdf"], "E203"), (["--scope", "some"], "E203")],
)
def test_export_command_rejects_bad_options(project: Workspace, args: list[str], code: str) -> None:
    result = runner.invoke(app, ["export", str(project.root), *args, "--lang", "en"])
    assert result.exit_code == 1 and code in result.output


def test_export_command_on_a_non_project(tmp_path: Path) -> None:
    result = runner.invoke(app, ["export", str(tmp_path / "nothing"), "--lang", "en"])
    assert result.exit_code == 1 and "E404" in result.output


def test_export_command_writes_the_prospero_draft(project: Workspace) -> None:
    pytest.importorskip("docx")
    result = runner.invoke(
        app, ["export", str(project.root), "--what", "prospero", "--lang", "en"]
    )
    assert result.exit_code == 0
    assert "PROSPERO protocol draft written" in result.output
