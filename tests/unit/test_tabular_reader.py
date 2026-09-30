"""Tests for the CSV/TSV/XLSX table reader (task T-M1-08, plan chapter 25.5)."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook

from crapai.errors import ImportFailed
from crapai.io.readers.detect import SourceFormat
from crapai.io.readers.tabular import (
    read_table,
    resolve_columns,
    unique_headers,
)

LEGACY = Path(__file__).resolve().parents[2] / "tests" / "legacy_runs"


def write_csv(path: Path, text: str, encoding: str = "utf-8") -> Path:
    path.write_bytes(text.encode(encoding))
    return path


BASIC = (
    (
        "Title;Abstract;Authors;Year;DOI;PMID;Keywords\n"
        "First study;An abstract, with comma;Doe, J; Roe, R;2020;10.1/a;32559806;x; y\n"
        "Second;;Solo A;n.d.;;0123;z\n"
    )
    .replace("Doe, J; Roe, R", '"Doe, J; Roe, R"')
    .replace("x; y", '"x; y"')
)


def test_semicolon_csv_is_mapped_by_alias(tmp_path: Path) -> None:
    result = read_table(write_csv(tmp_path / "a.csv", BASIC))
    assert result.format is SourceFormat.CSV
    assert result.options["delimiter"] == ";" and result.encoding == "utf-8-sig"
    first, second = result.records
    assert first.fields["title"] == "First study"
    assert first.fields["abstract"] == "An abstract, with comma"
    assert first.fields["authors"] == "Doe, J; Roe, R"
    assert first.fields["year"] == 2020 and first.fields["keywords"] == "x; y"
    assert "abstract" not in second.fields and "year" not in second.fields  # "n.d." -> no year
    assert [r.source_row for r in result.records] == [1, 2]
    assert result.column_map["abstract"] == "Abstract"


def test_pmid_and_doi_stay_text(tmp_path: Path) -> None:
    result = read_table(write_csv(tmp_path / "a.csv", BASIC))
    assert result.records[0].fields["pmid"] == "32559806"
    assert result.records[1].fields["pmid"] == "0123"  # leading zero kept
    spreadsheet = write_csv(tmp_path / "b.csv", "title,pmid\nT,32559806.0\n")
    assert read_table(spreadsheet).records[0].fields["pmid"] == "32559806"


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_delimiters_are_detected(tmp_path: Path, delimiter: str) -> None:
    text = delimiter.join(["title", "abstract"]) + "\n" + delimiter.join(["T1", "A1"]) + "\n"
    text += delimiter.join(["T2", "A2"]) + "\n"
    result = read_table(write_csv(tmp_path / "t.txt", text))
    assert result.options["delimiter"] == delimiter
    assert [r.fields["title"] for r in result.records] == ["T1", "T2"]


def test_forced_delimiter_and_encoding(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "x.csv", "title;abstract\nGröße;Ä\nB;C\n", encoding="cp1252")
    detected = read_table(path)
    assert detected.encoding == "cp1252" and detected.records[0].fields["title"] == "Größe"
    forced = read_table(path, encoding="cp1252", delimiter=";")
    assert forced.records[0].fields["abstract"] == "Ä"
    with pytest.raises(ImportFailed) as info:
        read_table(path, encoding="ascii")
    assert info.value.code == "E103"


def test_cp1252_and_utf8_files_are_both_read(tmp_path: Path) -> None:
    text = "title,abstract\nÄhnlichkeit – Größe,Ünder\nB,C\n"
    utf8 = read_table(write_csv(tmp_path / "u.csv", text, "utf-8"))
    cp = read_table(write_csv(tmp_path / "c.csv", text.replace("–", "-"), "cp1252"))
    assert utf8.encoding == "utf-8-sig" and cp.encoding == "cp1252"
    assert utf8.records[0].fields["title"] == "Ähnlichkeit – Größe"
    assert cp.records[0].fields["title"] == "Ähnlichkeit - Größe"


def test_legacy_run_csv_from_the_predecessor() -> None:
    """A real archived file: cp1252, semicolons, 99 rows, journal and journal_name columns."""
    path = sorted((LEGACY / "project-01_dhl").glob("*.csv"))[0]
    result = read_table(path)
    assert len(result.records) == 99 and result.encoding == "cp1252"
    assert result.options["delimiter"] == ";"
    assert result.column_map["journal"] == "journal"
    assert any("journal_name" in note for note in result.notes)  # ambiguity is reported
    with_abstract = [r for r in result.records if r.fields.get("abstract")]
    assert with_abstract and all(r.fields.get("title") for r in result.records[:5])
    assert "col_study_uid" in result.records[0].extra  # unmapped columns are kept


def test_headers_are_made_unique_and_blank_ones_named() -> None:
    assert unique_headers(["title", "Title ", "", "title", ""]) == [
        "title",
        "Title",
        "col_3",
        "title_2",
        "col_5",
    ]
    # the trimmed duplicate "Title" is a different string from "title", so only exact repeats count


def test_duplicate_header_is_numbered_in_a_file(tmp_path: Path) -> None:
    text = "title,abstract,title\nA,B,C\nD,E,F\n"
    result = read_table(write_csv(tmp_path / "d.csv", text))
    assert result.records[0].fields["title"] == "A"
    assert result.records[0].extra["col_title_2"] == "C"  # nothing lost


def test_alias_matching_ignores_case_space_and_underscore() -> None:
    headers = [
        "Article Title",
        "ABSTRACT_NOTE",
        "Author Full Names",
        "Publication Year",
        "PubMed ID",
    ]
    chosen, notes = resolve_columns(headers, {})
    assert chosen == {
        "title": "Article Title",
        "abstract": "ABSTRACT_NOTE",
        "authors": "Author Full Names",
        "year": "Publication Year",
        "pmid": "PubMed ID",
    }
    assert notes == []


def test_ambiguity_uses_alias_priority_and_reports(tmp_path: Path) -> None:
    text = "T1,Title,summary,abstract\nx,y,s,a\nx2,y2,s2,a2\n"
    result = read_table(write_csv(tmp_path / "m.csv", text))
    assert result.column_map["title"] == "Title"  # 'title' is listed before 't1'
    assert result.column_map["abstract"] == "abstract"
    assert any("--map title=" in n for n in result.notes)


def test_explicit_mapping_wins_and_is_returned(tmp_path: Path) -> None:
    text = "Name,Notes,Body\nPaper one,ignore,The real abstract\nPaper two,ignore,Another\n"
    path = write_csv(tmp_path / "e.csv", text)
    with pytest.raises(ImportFailed) as info:
        read_table(path)  # no title/abstract alias matches
    assert info.value.code == "E104" and "--map" in (info.value.hint or "")
    result = read_table(path, mapping={"title": "Name", "abstract": "Body"})
    assert result.column_map == {"title": "Name", "abstract": "Body"}
    assert result.records[0].fields["abstract"] == "The real abstract"
    assert result.records[0].extra == {"col_Notes": "ignore"}


def test_mapping_errors_name_the_available_columns(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "e.csv", "title,abstract\nA,B\nC,D\n")
    with pytest.raises(ImportFailed) as info:
        read_table(path, mapping={"abstract": "Zusammenfassung"})
    assert info.value.code == "E104" and "title, abstract" in (info.value.hint or "")
    with pytest.raises(ImportFailed) as info2:
        read_table(path, mapping={"colour": "title"})
    assert info2.value.code == "E104"


def test_missing_abstract_column_is_a_note_not_an_error(tmp_path: Path) -> None:
    result = read_table(write_csv(tmp_path / "t.csv", "title,year\nA,2020\nB,2021\n"))
    assert any(note.startswith("no abstract column found") for note in result.notes)
    assert "abstract" not in result.column_map


def test_header_only_and_ragged_rows(tmp_path: Path) -> None:
    with pytest.raises(ImportFailed) as info:
        read_table(write_csv(tmp_path / "h.csv", "title,abstract,year\n"))
    assert info.value.code in {"E101", "E102"}  # a lone header row is not even a table
    body = "".join(f"T{i},A{i}\n" for i in range(8))  # 8 regular rows keep the sniffer sure
    ragged = write_csv(tmp_path / "r.csv", "title,abstract\nshort\n" + body + "D,E,extra cell\n")
    result = read_table(ragged)
    assert result.records[0].fields["title"] == "short"
    assert "abstract" not in result.records[0].fields
    assert result.records[-1].fields["title"] == "D"
    assert result.records[-1].extra["col_overflow"] == ["extra cell"]
    assert result.records[-1].notes


def test_non_table_formats_are_refused() -> None:
    with pytest.raises(ImportFailed) as info:
        read_table(Path(__file__).resolve().parents[2] / "tests" / "data" / "example_AB_nr4.ris")
    assert info.value.code == "E101"


# --- Excel ---


def make_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Records"
    sheet.append(["Title", "Abstract", "PMID", "Year", "DOI", "Added"])
    sheet.append(["Paper A", "Abstract A", 32559806, 2020, "10.1/a", dt.datetime(2021, 1, 4)])
    sheet.append(["Paper B", None, 32559807.0, 2019.0, "10.1/b", dt.datetime(2021, 1, 4, 13, 30)])
    sheet.append([None, None, None, None, None, None])  # empty row is skipped
    hidden = workbook.create_sheet("Old")
    hidden.sheet_state = "hidden"
    hidden.append(["Title", "Abstract"])
    hidden.append(["Hidden paper", "x"])
    other = workbook.create_sheet("More")
    other.append(["Title", "Abstract"])
    other.append(["Other paper", "y"])
    workbook.save(path)
    return path


def test_xlsx_first_visible_sheet_with_text_ids_and_dates(tmp_path: Path) -> None:
    result = read_table(make_workbook(tmp_path / "book.xlsx"))
    assert result.format is SourceFormat.XLSX and result.encoding is None
    assert result.options["sheet"] == "Records"
    assert [r.fields["title"] for r in result.records] == ["Paper A", "Paper B"]
    assert result.records[0].fields["pmid"] == "32559806"  # no 3.2559806E+07, no ".0"
    assert result.records[1].fields["pmid"] == "32559807"
    assert result.records[1].fields["year"] == 2019
    assert result.records[0].extra["col_Added"] == "2021-01-04"
    assert result.records[1].extra["col_Added"] == "2021-01-04T13:30:00"
    assert any("2 sheets" in note for note in result.notes)  # hidden sheet is not counted


def test_xlsx_named_sheet_and_unknown_sheet(tmp_path: Path) -> None:
    path = make_workbook(tmp_path / "book.xlsx")
    assert read_table(path, sheet="More").records[0].fields["title"] == "Other paper"
    with pytest.raises(ImportFailed) as info:
        read_table(path, sheet="Old")  # hidden sheets are ignored
    assert info.value.code == "E101" and "Records" in (info.value.hint or "")


def test_xlsx_formulas_are_read_as_values_only(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Title", "Abstract"])
    sheet.append(['="A"&"B"', "abstract"])
    sheet.append(["plain", "abstract 2"])
    path = tmp_path / "f.xlsx"
    workbook.save(path)
    result = read_table(path)
    # openpyxl has no cached value for a formula it wrote itself: an empty cell, never the formula
    assert "=" not in result.records[0].fields.get("title", "")


def test_corrupt_xlsx_is_reported(tmp_path: Path) -> None:
    good = make_workbook(tmp_path / "ok.xlsx")
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(good.read_bytes()[:200])
    with pytest.raises(ImportFailed) as info:
        read_table(broken)
    assert info.value.code == "E101"
