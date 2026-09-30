"""Edge cases and defensive branches found with a branch-coverage audit (measured with coverage.py).

Each test names the behaviour it protects. They complement the per-module test files; where a
behaviour belongs to one module only, the module is named in the section heading.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import logging
import logging.handlers
import runpy
import sys
import types
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from pydantic import ValidationError
from typer.testing import CliRunner

from crapai import cli
from crapai.cli import app
from crapai.config import loader
from crapai.config.models import ProjectInfo
from crapai.errors import ConfigError, ImportFailed, SaraError, StorageError
from crapai.i18n.loader import I18n
from crapai.i18n.messages import Messages, resolve_language
from crapai.io.import_log import ImportLogEntry, append_entry, read_entries
from crapai.io.readers import detect, tabular
from crapai.io.readers.bibtex import parse_bibtex_text
from crapai.io.readers.detect import SourceFormat, detect_format
from crapai.io.readers.nbib import parse_nbib_text
from crapai.io.readers.ris import parse_ris_text
from crapai.io.readers.tabular import read_table
from crapai.io.records_store import RECORD_COLUMNS, Record, read_records, write_records
from crapai.project import lock as lock_module
from crapai.project.lock import ProjectLock
from crapai.project.workspace import Workspace
from crapai.services.importing import ImportRequest, import_source
from crapai.services.project import create_project

DATA = Path(__file__).resolve().parents[2] / "tests" / "data"
runner = CliRunner()


def demo(tmp_path: Path, name: str = "p") -> Path:
    return create_project(tmp_path / name, template="demo").root


# --- cli -------------------------------------------------------------------------------------


def test_the_project_log_follows_the_project_and_is_not_attached_twice(tmp_path: Path) -> None:
    first, second = demo(tmp_path, "one"), demo(tmp_path, "two")
    ris = DATA / "example_AB_nr4.ris"
    runner.invoke(app, ["import", str(first), str(ris)])
    runner.invoke(app, ["dedup", str(first)])  # same project again: no second handler
    handlers = [
        h
        for h in logging.getLogger("crapai").handlers
        if isinstance(h, logging.handlers.RotatingFileHandler)
    ]
    assert len(handlers) == 1 and handlers[0].baseFilename.startswith(str(first.resolve()))
    runner.invoke(app, ["dedup", str(second)])  # another project: the handler moves
    handlers = [
        h
        for h in logging.getLogger("crapai").handlers
        if isinstance(h, logging.handlers.RotatingFileHandler)
    ]
    assert len(handlers) == 1 and handlers[0].baseFilename.startswith(str(second.resolve()))


def test_json_mode_prints_partial_results_before_an_error(tmp_path: Path) -> None:
    project = demo(tmp_path)
    missing = tmp_path / "missing.ris"
    args = ["import", str(project), str(missing), "--json", "--lang", "en"]
    result = CliRunner().invoke(app, args)
    assert result.exit_code == 1
    document = json.loads(result.stdout)  # stdout stays one valid JSON document
    assert document["imported"] == [] and document["error"]["code"] == "E101"
    assert "Error E101" in result.stderr


def test_first_file_imported_then_second_fails_reports_the_first_in_json(tmp_path: Path) -> None:
    project = demo(tmp_path)
    good, bad = DATA / "example_AB_nr4.ris", tmp_path / "nope.ris"
    result = CliRunner().invoke(app, ["import", str(project), str(good), str(bad), "--json"])
    assert result.exit_code == 1
    imported = json.loads(result.stdout)["imported"]
    assert [entry["records"] for entry in imported] == [3]


def test_excel_import_has_no_encoding_line(tmp_path: Path) -> None:
    project = demo(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Title", "Abstract"])
    sheet.append(["Paper one", "Text one"])
    sheet.append(["Paper two", "Text two"])
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    result = runner.invoke(app, ["import", str(project), str(path), "--lang", "en"])
    assert result.exit_code == 0, result.output
    assert "Text encoding" not in result.output and "Format: xlsx" in result.output


def test_status_tells_when_another_process_holds_the_project(tmp_path: Path) -> None:
    project = demo(tmp_path)
    with Workspace(project).lock():
        result = runner.invoke(app, ["status", str(project), "--lang", "en"])
    assert "in use by another process" in result.output


def test_main_switches_a_legacy_console_to_utf8_and_runs_the_app(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[str] = []

    class FakeStdout:
        encoding = "cp1252"

        def reconfigure(self, *, encoding: str, errors: str) -> None:
            calls.append(f"{encoding}/{errors}")

    monkeypatch.setattr(sys, "stdout", FakeStdout())
    monkeypatch.setattr(cli, "app", lambda: calls.append("app"))
    cli.main()
    assert calls == ["utf-8/replace", "app"]


def test_main_leaves_a_utf8_console_alone(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []
    fake = types.SimpleNamespace(encoding="utf-8", reconfigure=lambda **kw: calls.append("bad"))
    monkeypatch.setattr(sys, "stdout", fake)
    monkeypatch.setattr(cli, "app", lambda: calls.append("app"))
    cli.main()
    assert calls == ["app"]


def test_running_the_module_as_a_script_shows_the_version(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(sys, "argv", ["crapai", "--version"])
    monkeypatch.delitem(sys.modules, "crapai.cli", raising=False)
    with pytest.raises(SystemExit) as info:
        runpy.run_module("crapai.cli", run_name="__main__")
    assert info.value.code == 0
    assert "Critical Apprais.AI" in capsys.readouterr().out


# --- configuration ---------------------------------------------------------------------------


def test_default_user_config_lives_under_dot_config_crapai(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    assert loader.default_user_config_path() == tmp_path / ".config" / "crapai" / "config.yaml"


def test_an_override_that_is_not_valid_yaml_stays_a_string() -> None:
    assert loader.env_overrides({"CRAPAI_LLM__MODEL": "[1, 2"}) == {"llm": {"model": "[1, 2"}}
    assert loader.parse_cli_overrides(["llm.model={unclosed"]) == {"llm": {"model": "{unclosed"}}


def test_blank_project_title_is_rejected() -> None:
    with pytest.raises(ValidationError) as info:
        ProjectInfo(title="   ")
    assert "must not be blank" in str(info.value)


def test_user_config_that_is_not_a_mapping_is_a_config_error(tmp_path: Path) -> None:
    project = demo(tmp_path)
    user = tmp_path / "user.yaml"
    user.write_text("- not\n- a mapping\n", encoding="utf-8")
    with pytest.raises(ConfigError) as info:
        loader.resolve_config(project / "project.yaml", environ={}, user_config_path=user)
    assert info.value.code == "E203" and "mapping" in info.value.user_message


# --- import log ------------------------------------------------------------------------------


def test_blank_lines_inside_the_import_log_are_ignored(tmp_path: Path) -> None:
    log = tmp_path / "log.jsonl"
    entry = ImportLogEntry(
        timestamp=dt.datetime(2026, 9, 30, tzinfo=dt.UTC),
        source_file="a.ris",
        sha256="a" * 64,
        source_label="A",
        format="ris",
        records=1,
        abstracts=1,
    )
    append_entry(log, entry)
    with open(log, "a", encoding="utf-8") as handle:
        handle.write("\n   \n")
    append_entry(log, entry)
    assert len(read_entries(log)) == 2


# --- readers ---------------------------------------------------------------------------------


def test_bibtex_entry_without_cite_key_and_with_truncated_value() -> None:
    result = parse_bibtex_text("@misc{ title = {No key here}, year = 2020 }\n@misc{k, title = }\n")
    assert result.records[0].fields["title"] == "No key here"
    assert "bib_key" not in result.records[0].extra
    assert result.records[1].extra["bib_key"] == "k" and "title" not in result.records[1].fields


def test_bibtex_unterminated_quoted_value_is_read_to_the_end() -> None:
    result = parse_bibtex_text('@misc{k, title = "never closed')
    assert result.records[0].fields["title"] == "never closed"


def test_mixed_markers_with_ris_first_are_read_as_ris_with_lower_confidence(
    tmp_path: Path,
) -> None:
    path = tmp_path / "m.txt"
    path.write_text("TY  - JOUR\nTI  - A\nER  - \n\nPMID- 9\nTI  - B\n", encoding="utf-8")
    found = detect_format(path)
    assert found.format is SourceFormat.RIS and found.confidence < 0.9
    assert "before any MEDLINE" in found.reason


def test_a_csv_parser_error_while_sniffing_means_not_a_table(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken_reader(*args: Any, **kwargs: Any) -> Any:
        raise csv.Error("simulated parser error")

    monkeypatch.setattr(detect.csv, "reader", broken_reader)
    path = tmp_path / "t.txt"
    path.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    with pytest.raises(ImportFailed) as info:
        detect_format(path)
    assert info.value.code == "E101"


def test_nbib_keeps_a_second_different_doi_marker_as_other_identifier() -> None:
    text = "PMID- 1\nTI  - T\nLID - 10.1/a [doi]\nAID - 10.1/b [doi]\nAID - 10.1/a [doi]\n"
    record = parse_nbib_text(text).records[0]
    assert record.fields["doi"] == "10.1/a"
    assert record.extra["nbib_AID"] == "10.1/b [doi]"  # the conflicting DOI is not lost


def test_ris_header_note_is_shortened() -> None:
    long_header = "Provider: " + "x" * 400
    text = f"{long_header}\nTY  - JOUR\nTI  - T\nER  - \n"
    (record,) = parse_ris_text(text).records
    (note,) = record.notes
    assert note.endswith("...") and len(note) < 260


# --- table reader ----------------------------------------------------------------------------


def test_cell_text_conversions() -> None:
    conv = tabular._cell_text
    assert conv(None) == "" and conv(True) == "true" and conv(False) == "false"
    assert conv(3.0) == "3" and conv(2.5) == "2.5" and conv(7) == "7"
    assert conv(dt.date(2021, 3, 4)) == "2021-03-04"
    assert conv(dt.datetime(2021, 3, 4)) == "2021-03-04"
    assert conv(dt.datetime(2021, 3, 4, 5, 6)) == "2021-03-04T05:06:00"
    assert conv("  padded  ") == "padded"


def test_undecodable_bytes_in_a_csv_are_counted_in_the_notes(tmp_path: Path) -> None:
    path = tmp_path / "in.csv"
    path.write_text("title,abstract\nA � B,x\nC,y\n", encoding="utf-8")
    result = read_table(path)
    assert any("1 replacement character" in note for note in result.notes)


def test_zip_that_only_looks_like_a_workbook_is_reported(tmp_path: Path) -> None:
    import zipfile

    path = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "this is not a workbook")
    with pytest.raises(ImportFailed) as info:
        read_table(path)
    assert info.value.code == "E101" and "Excel workbook" in info.value.user_message


def test_workbook_without_a_visible_sheet_is_reported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """openpyxl refuses to write such a file, so the workbook object is simulated."""
    real = Workbook()
    first = real.active
    assert first is not None
    first.append(["Title", "Abstract"])
    path = tmp_path / "book.xlsx"
    real.save(path)  # a real file so that the format detection accepts it

    class HiddenBook:
        worksheets = [types.SimpleNamespace(sheet_state="hidden", title="Secret")]

        def close(self) -> None:
            pass

    monkeypatch.setattr("openpyxl.load_workbook", lambda *args, **kwargs: HiddenBook())
    with pytest.raises(ImportFailed) as info:
        read_table(path)
    assert info.value.code == "E102" and "no visible worksheet" in info.value.user_message


def test_header_only_workbook_has_no_data_rows(tmp_path: Path) -> None:
    header_only = Workbook()
    sheet = header_only.active
    assert sheet is not None
    sheet.append(["Title", "Abstract"])
    path = tmp_path / "header.xlsx"
    header_only.save(path)
    with pytest.raises(ImportFailed) as info:
        read_table(path)
    assert info.value.code == "E102"


def test_a_large_sheet_gets_a_note(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(tabular, "LARGE_SHEET_ROWS", 2)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Title", "Abstract"])
    for index in range(3):
        sheet.append([f"T{index}", f"A{index}"])
    path = tmp_path / "big.xlsx"
    workbook.save(path)
    assert any("large sheet (4 rows)" in note for note in read_table(path).notes)


# --- records file ----------------------------------------------------------------------------


def _record(**overrides: Any) -> Record:
    values: dict[str, Any] = {
        "study_uid": "u1",
        "source_label": "A",
        "source_file": "a.ris",
        "source_row": 1,
        "source_format": "ris",
    }
    values.update(overrides)
    return Record(**values)


def test_empty_boolean_cells_mean_false_and_blank_lines_are_skipped(tmp_path: Path) -> None:
    path = tmp_path / "records.csv"
    write_records(path, [_record(), _record(study_uid="u2", source_row=2)])
    lines = path.read_text(encoding="utf-8").split("\n")
    cells = next(csv.reader([lines[1]]))
    for name in ("is_retracted", "has_abstract"):
        cells[RECORD_COLUMNS.index(name)] = ""
    lines[1] = ",".join(cells)
    lines.insert(2, "")  # an empty line between the rows
    path.write_text("\n".join(lines), encoding="utf-8", newline="")
    loaded = read_records(path)
    assert [r.study_uid for r in loaded] == ["u1", "u2"]
    assert loaded[0].is_retracted is False and loaded[0].has_abstract is False


# --- import service --------------------------------------------------------------------------


def test_a_bad_copy_into_sources_is_detected_and_removed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workspace = create_project(tmp_path / "p", template="demo")

    def corrupting_copy(src: Any, dst: Any) -> None:
        Path(dst).write_bytes(b"not the same bytes")

    monkeypatch.setattr("crapai.services.importing.shutil.copyfile", corrupting_copy)
    with pytest.raises(StorageError) as info:
        import_source(workspace, ImportRequest(DATA / "example_AB_nr4.ris"))
    assert info.value.code == "E401" and "checksum" in info.value.user_message
    assert not list(workspace.sources_dir.iterdir())  # the damaged copy is gone
    assert not workspace.records_csv.exists() and not workspace.import_log.exists()


# --- project lock ----------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raised", "expected"),
    [(None, True), (ProcessLookupError(), False), (PermissionError(), True)],
)
def test_pid_alive_on_posix(
    monkeypatch: pytest.MonkeyPatch, raised: Exception | None, expected: bool
) -> None:
    def fake_kill(pid: int, signal: int) -> None:
        assert signal == 0  # only an existence check, never a real signal
        if raised is not None:
            raise raised

    monkeypatch.setattr(lock_module.sys, "platform", "linux")
    monkeypatch.setattr(lock_module.os, "kill", fake_kill)
    assert lock_module.pid_alive(4242) is expected


@pytest.mark.skipif(sys.platform != "win32", reason="Win32 API")
def test_windows_pid_check_handles_access_denied_and_query_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import ctypes

    class Function:
        """A Win32 function stand-in that accepts the argtypes/restype declarations."""

        def __init__(self, result: int) -> None:
            self.result = result
            self.argtypes: Any = None
            self.restype: Any = None

        def __call__(self, *args: Any) -> int:
            return self.result

    fake = types.SimpleNamespace(
        OpenProcess=Function(0),  # no handle
        GetExitCodeProcess=Function(0),  # the query fails
        CloseHandle=Function(1),
    )
    monkeypatch.setattr(ctypes, "WinDLL", lambda *a, **k: fake, raising=False)
    monkeypatch.setattr(ctypes, "get_last_error", lambda: 5, raising=False)
    assert lock_module.pid_alive(1234) is True  # access denied: the process exists
    monkeypatch.setattr(ctypes, "get_last_error", lambda: 87, raising=False)
    assert lock_module.pid_alive(1234) is False  # invalid parameter: no such process
    fake.OpenProcess.result = 99
    assert lock_module.pid_alive(1234) is True  # handle opened but query failed: assume alive
    assert fake.OpenProcess.restype is not None  # 64-bit handles are declared, not truncated


def test_unreadable_lock_file_is_assumed_to_be_in_use(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "lock"
    path.mkdir()  # reading a directory as a file raises an OSError
    monkeypatch.setattr(lock_module, "READ_RETRY_DELAY_S", 0.0)
    with caplog.at_level("WARNING"):
        info = ProjectLock(path).inspect()
    # an unknown state must never allow a takeover of a lock that may be alive
    assert info is not None and not info.stale and info.pid == 0
    assert "assuming it is in use" in caplog.text


def test_a_short_read_error_is_retried_and_then_succeeds(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    lock = ProjectLock(tmp_path / "lock", pid=7)
    lock.acquire()
    real = Path.read_text
    calls = {"n": 0}

    def flaky(self: Path, *args: Any, **kwargs: Any) -> str:
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError("held by a virus scanner")
        return real(self, *args, **kwargs)

    monkeypatch.setattr(lock_module, "READ_RETRY_DELAY_S", 0.0)
    monkeypatch.setattr(Path, "read_text", flaky)
    info = lock.inspect()
    assert info is not None and info.pid == 7 and not info.stale and calls["n"] == 3


def test_lock_file_that_is_not_utf8_is_stale(tmp_path: Path) -> None:
    path = tmp_path / "lock"
    path.write_bytes(b"\xff\xfe\x00garbage")
    info = ProjectLock(path).inspect()
    assert info is not None and info.stale


def test_timestamps_without_time_zone_are_read_as_utc(tmp_path: Path) -> None:
    path = tmp_path / "lock"
    path.write_text(
        '{"pid": 1, "token": "t", "started": "2026-01-01T00:00:00",'
        ' "heartbeat": "2026-01-01T00:00:00"}',
        encoding="utf-8",
    )
    info = ProjectLock(path, is_alive=lambda pid: False).inspect()  # used to raise TypeError
    assert info is not None and info.stale and info.heartbeat.tzinfo is not None


def test_remove_stale_leaves_a_live_lock_alone(tmp_path: Path) -> None:
    path = tmp_path / "lock"
    ProjectLock(path, pid=1).acquire()
    assert ProjectLock(path, pid=2).remove_stale() is False
    assert path.exists()


def test_remove_stale_removes_a_dead_owners_lock(tmp_path: Path) -> None:
    path = tmp_path / "lock"
    ProjectLock(path, pid=1).acquire()
    later = ProjectLock(
        path,
        pid=2,
        is_alive=lambda pid: False,
        now=lambda: dt.datetime.now(dt.UTC) + dt.timedelta(seconds=120),
    )
    assert later.remove_stale() is True
    assert not path.exists() and list(tmp_path.iterdir()) == []  # no parked file is left behind


def test_remove_stale_puts_back_a_lock_that_changed_hands(tmp_path: Path) -> None:
    path = tmp_path / "lock"
    ProjectLock(path, pid=1).acquire()
    stale_view = ProjectLock(
        path,
        pid=2,
        is_alive=lambda pid: False,
        now=lambda: dt.datetime.now(dt.UTC) + dt.timedelta(seconds=120),
    ).inspect()
    assert stale_view is not None and stale_view.stale
    path.unlink()
    fresh = ProjectLock(path, pid=3)
    fresh.acquire()  # another process re-created the lock after our inspection
    other = ProjectLock(path, pid=2, is_alive=lambda pid: False)
    assert other.remove_stale(stale_view) is False
    assert fresh.held  # the fresh lock was put back


def test_lock_errors_become_storage_errors(tmp_path: Path) -> None:
    blocker = tmp_path / "file"
    blocker.write_text("x", encoding="utf-8")
    with pytest.raises(StorageError) as info:
        ProjectLock(blocker / "sub" / "lock").acquire()  # the folder cannot be created
    assert info.value.code == "E401"
    directory = tmp_path / "dir"
    directory.mkdir()
    with pytest.raises(StorageError) as again:
        ProjectLock(directory).acquire()  # the lock path is a directory
    assert again.value.code in {"E401", "E402"}


def test_lock_that_vanishes_between_two_checks_is_still_acquired(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    lock = ProjectLock(tmp_path / "lock")
    attempts = iter([False, True])
    monkeypatch.setattr(lock, "_try_create", lambda: next(attempts))
    monkeypatch.setattr(lock, "inspect", lambda: None)  # released in the meantime
    lock.acquire()  # first create fails, inspect sees nothing, second create works


def test_takeover_that_loses_the_race_is_refused(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "lock"
    old = ProjectLock(path, pid=1, is_alive=lambda pid: False)
    old.acquire()
    later = ProjectLock(
        path,
        pid=2,
        is_alive=lambda pid: False,
        now=lambda: dt.datetime.now(dt.UTC) + dt.timedelta(seconds=120),
    )
    real_create = later._try_create
    calls = {"count": 0}

    def losing_create() -> bool:
        calls["count"] += 1
        return False if calls["count"] >= 2 else real_create()

    monkeypatch.setattr(later, "_try_create", losing_create)
    with pytest.raises(StorageError) as info:
        later.acquire(take_over_stale=True)
    assert info.value.code == "E402"


# --- messages and texts ----------------------------------------------------------------------


def test_project_language_that_is_not_supported_falls_through_to_the_environment(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("LC_ALL", raising=False)
    monkeypatch.setenv("LANG", "de_CH.UTF-8")
    project = tmp_path / "project.yaml"
    project.write_text("project:\n  title: t\n  language: fr\n", encoding="utf-8")
    assert resolve_language(None, project) == "de"


def test_error_report_without_message_or_hint() -> None:
    lines = Messages("en").error_lines(SaraError("", code="E999"))
    assert lines[0].startswith("Error E999")
    assert not any(line.startswith("Details") for line in lines)  # empty message: no empty line
    unknown = Messages("en").error_lines(SaraError("", code="E777", hint="do this"))
    assert unknown[0].startswith("Error E777")  # the real code is shown ...
    assert unknown[-1].startswith("What to do")  # ... with the generic advice


def test_i18n_loader_with_own_folder_and_damaged_files(tmp_path: Path) -> None:
    texts = tmp_path / "texts"
    texts.mkdir()
    (texts / "en.yaml").write_text("greeting: Hello\nempty: ''\nitems: []\n", encoding="utf-8")
    (texts / "de.yaml").write_text("- this is a list, not a mapping\n", encoding="utf-8")
    (texts / "fr.yaml").write_text("greeting: [unclosed\n", encoding="utf-8")
    for language in ("de", "fr", "xx"):
        i18n = I18n(base_dir=str(tmp_path))
        i18n.load(language)
        assert i18n.t("greeting") == "Hello"  # the damaged overlay is ignored, English stays
    i18n = I18n(base_dir=str(tmp_path))
    i18n.load("en")
    assert i18n.validate(["greeting", "empty", "items", "missing.key"]) == [
        "empty",
        "items",
        "missing.key",
    ]
    assert set(i18n.resolved_paths) == {"en"}
    assert i18n.t("greeting.deeper") == ""  # walking into a string is not an error
